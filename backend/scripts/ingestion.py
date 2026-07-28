import os
import sys
import re
import docx

# Add parent directory (backend) to sys.path so config can be imported
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from langchain_core.documents import Document
from langchain_chroma import Chroma
from config import logger
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

DB_DIR = "./chroma_db"

embeddings = OllamaEmbeddings(
    model="bge-m3:latest",
    base_url=os.getenv("OLLAMA_HOST", "http://localhost:11434")
)

def ingest_data():
    file_path = "../data/FAQ_Demandes_et_informations_with_links.docx"
    
    if not os.path.exists(file_path):
        logger.error(f"File not found: {file_path}")
        return

    doc = docx.Document(file_path)
    
    CATEGORIES = ["Commerciale", "Technique"]
    SUBCATEGORIES = {
        "Commerciale": ["Branchement", "Abonnement", "Résiliation", "Tarification", "Solutions de paiement", "Services digitaux", "Service SMS", "Demande d'attestations", "Réseau Commercial"],
        "Technique": ["Branchement", "Assainissement", "Eau", "Électricité", "Coupure des fournitures"]
    }
    
    docs = []
    
    current_category = "commerciale"
    current_subcategory = "branchement"
    current_question = None
    current_answer = ""
    
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=100,
        length_function=len,
    )
    
    def flush_qa():
        nonlocal current_answer, current_question, docs, current_category, current_subcategory
    
        if current_question and len(current_answer.strip()) >= 20:
            category = current_category
            subcategory = current_subcategory
            question = current_question
            
            chunks = text_splitter.split_text(current_answer.strip())
            
            for chunk in chunks:
                docs.append(Document(
                    page_content=chunk,
                    metadata={
                        "category": category,
                        "subcategory": subcategory,
                        "question": question,
                        "source": "faq_docx"
                    }
                ))
        current_answer = ""
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
            
        # Standardize bullet points to markdown
        text = text.replace("›", "-").replace("•", "-")
            
        text_lower = text.lower()
        
        # 1. Detect Category
        found_cat = False
        if "commerciale" in text_lower and len(text_lower) < 50:
            flush_qa()
            current_category = "commerciale"
            current_question = None
            found_cat = True
        elif "technique" in text_lower and len(text_lower) < 50:
            flush_qa()
            current_category = "technique"
            current_question = None
            found_cat = True
        if found_cat:
            continue
            
        # 2. Detect Question
        q_match = re.match(r"^(Q|Question)\s*[:：]\s*", text, re.IGNORECASE)
        if q_match:
            flush_qa()
            current_question = text[q_match.end():].strip().lower()
            continue

        # 3. Detect Subcategory
        found_sub = False
        for cat in CATEGORIES:
            for sub in SUBCATEGORIES[cat]:
                if sub.lower() in text_lower and len(text_lower) < 50:
                    flush_qa()
                    current_category = cat.strip().lower()
                    current_subcategory = sub.strip().lower()
                    current_question = None
                    found_sub = True
                    break
            if found_sub:
                break
        if found_sub:
            continue
            
        # 4. Detect Answer
        r_match = re.match(r"^(R|Réponse)\s*[:：]\s*", text, re.IGNORECASE)
        if r_match:
            answer_part = text[r_match.end():].strip()
            if answer_part:
                current_answer += answer_part + "\n"
            continue
            
        # 5. Multiline Answer Accumulation
        if current_question:
            current_answer += text + "\n"

    # Flush any remaining Q/A
    flush_qa()
    
    if not docs:
        logger.warning("No structured documents parsed.")
        return
        
    logger.info(f"Storing {len(docs)} chunks in ChromaDB...")
    vectorstore = Chroma.from_documents(
        documents=docs,
        embedding=embeddings,
        persist_directory=DB_DIR
    )

    logger.info(f"Ingestion complete at {DB_DIR}")

if __name__ == "__main__":
    ingest_data()
